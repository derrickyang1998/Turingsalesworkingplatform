#!/usr/bin/env python3
"""Render one sealed TuringMarket customer report into a PPTX artifact."""

import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
SKY = RGBColor(219, 234, 254)
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
BORDER = RGBColor(226, 232, 240)
REQUIRED_SECTIONS = (
    "project_overview",
    "data_summary",
    "eligible_comparisons",
    "key_indicators",
    "excellent_cases",
    "data_limits_and_risks",
    "optimization_and_next_cycle",
)


def abort(message):
    raise ValueError(message)


def as_mapping(value, label):
    if not isinstance(value, dict):
        abort(f"{label} must be an object")
    return value


def as_list(value):
    return value if isinstance(value, list) else []


def text(value, fallback="未提供"):
    if value is None:
        return fallback
    rendered = str(value).strip()
    return rendered if rendered else fallback


def number(value, fallback="--"):
    if isinstance(value, bool) or value is None:
        return fallback
    if isinstance(value, (int, float)):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return text(value, fallback)


def metric_value(metric):
    source = metric if isinstance(metric, dict) else {}
    return text(source.get("display_value"), number(source.get("value")))


def load_report(input_path):
    if not os.path.isabs(input_path):
        abort("input path must be absolute")
    with open(input_path, "r", encoding="utf-8") as source:
        report = json.load(source)
    report = as_mapping(report, "report")
    if report.get("contract_version") != "customer_safe_v1":
        abort("customer report contract is invalid")
    if report.get("redaction_policy_version") != "customer-safe-v1":
        abort("customer report redaction policy is invalid")
    if report.get("recipient_profile") != "customer" or report.get("status") != "sealed":
        abort("customer report is not sealed for delivery")
    sections = as_mapping(report.get("sections"), "sections")
    if any(key not in sections for key in REQUIRED_SECTIONS):
        abort("customer report sections are incomplete")
    return report, sections


def add_textbox(slide, left, top, width, height, value, size=18, color=TEXT, bold=False,
                alignment=PP_ALIGN.LEFT, font_name="Aptos"):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text(value, "")
    paragraph.alignment = alignment
    paragraph.font.name = font_name
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def add_base_slide(presentation, title, subtitle=None):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = WHITE
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.12))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    add_textbox(slide, Inches(0.72), Inches(0.42), Inches(11.9), Inches(0.45), title, 27, NAVY, True)
    if subtitle:
        add_textbox(slide, Inches(0.74), Inches(0.93), Inches(11.8), Inches(0.32), subtitle, 10.5, MUTED)
    footer = add_textbox(slide, Inches(0.74), Inches(7.07), Inches(11.8), Inches(0.2), "TuringMarket · 客户版项目复盘", 8.5, MUTED)
    footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return slide


def add_card(slide, left, top, width, height, label, value, detail=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(248, 250, 252)
    card.line.color.rgb = BORDER
    add_textbox(slide, left + Inches(0.2), top + Inches(0.16), width - Inches(0.4), Inches(0.25), label, 10, MUTED)
    add_textbox(slide, left + Inches(0.2), top + Inches(0.48), width - Inches(0.4), Inches(0.42), value, 19, NAVY, True)
    if detail:
        add_textbox(slide, left + Inches(0.2), top + Inches(0.97), width - Inches(0.4), height - Inches(1.1), detail, 9.5, MUTED)


def add_bullet_list(slide, items, top=Inches(1.55), max_items=6):
    values = [text(item, "") for item in as_list(items) if text(item, "")]
    if not values:
        values = ["当前数据范围内暂无可公开的结论。"]
    box = slide.shapes.add_textbox(Inches(0.92), top, Inches(11.45), Inches(4.95))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(values[:max_items]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(17)
        paragraph.font.color.rgb = TEXT
        paragraph.space_after = Pt(13)


def add_comparison_rows(slide, groups):
    row_top = Inches(1.56)
    labels = [("平台", as_list(groups.get("platforms"))), ("产品", as_list(groups.get("products")))]
    rows = []
    for group_label, items in labels:
        for item in items:
            if not isinstance(item, dict):
                continue
            rows.append((
                group_label,
                text(item.get("label")),
                text(item.get("content_count"), "0") + " 条内容",
                metric_value(item.get("selected_metric")),
            ))
    if not rows:
        add_bullet_list(slide, [text(groups.get("reason"), "暂无可公开的比较维度。")])
        return
    headers = ("维度", "分组", "内容量", "表现")
    widths = (Inches(1.2), Inches(4.35), Inches(2.3), Inches(3.7))
    left = Inches(0.95)
    for index, header in enumerate(headers):
        add_textbox(slide, left, row_top, widths[index], Inches(0.36), header, 11, MUTED, True)
        left += widths[index]
    for row_index, row in enumerate(rows[:8]):
        y = row_top + Inches(0.47 + row_index * 0.52)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.88), y - Inches(0.06), Inches(11.55), Inches(0.44))
        band.fill.solid()
        band.fill.fore_color.rgb = RGBColor(248, 250, 252) if row_index % 2 == 0 else WHITE
        band.line.color.rgb = BORDER
        left = Inches(0.98)
        for column, value in enumerate(row):
            add_textbox(slide, left, y, widths[column], Inches(0.3), value, 11.5, TEXT, column == 3)
            left += widths[column]


def build_presentation(report, sections):
    overview = as_mapping(sections["project_overview"], "project_overview")
    summary = as_mapping(sections["data_summary"], "data_summary")
    indicators = as_mapping(sections["key_indicators"], "key_indicators")
    limits = as_mapping(sections["data_limits_and_risks"], "data_limits_and_risks")
    optimization = as_mapping(sections["optimization_and_next_cycle"], "optimization_and_next_cycle")
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = cover.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    cover_accent = cover.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.16), SLIDE_HEIGHT
    )
    cover_accent.fill.solid()
    cover_accent.fill.fore_color.rgb = BLUE
    cover_accent.line.fill.background()
    add_textbox(cover, Inches(0.95), Inches(2.08), Inches(11.25), Inches(0.75), text(report.get("title"), "项目复盘"), 34, WHITE, True)
    add_textbox(cover, Inches(0.98), Inches(3.0), Inches(11.1), Inches(0.36), text(overview.get("campaign_name"), "推广项目"), 16, RGBColor(191, 219, 254))
    add_textbox(cover, Inches(0.98), Inches(5.87), Inches(11.1), Inches(0.32), "客户版复盘 · 基于已封存数据快照", 11, RGBColor(203, 213, 225))

    project = add_base_slide(presentation, "项目概况", text(overview.get("campaign_name"), "推广项目"))
    cards = (
        ("内容数量", text(overview.get("content_count"), "0") + " 条", None),
        ("平台", " · ".join(as_list(overview.get("platform_mix"))) or "未提供", None),
        ("观测窗口", text(overview.get("observation_window", {}).get("label") if isinstance(overview.get("observation_window"), dict) else ""), None),
    )
    for index, (label, value, detail) in enumerate(cards):
        add_card(project, Inches(0.93 + index * 4.0), Inches(1.6), Inches(3.7), Inches(1.42), label, value, detail)
    coverage = []
    for item in as_list(overview.get("data_coverage")):
        if isinstance(item, dict):
            coverage.append(text(item.get("metric")) + "：" + text(item.get("available_records"), "0") + " / " + text(item.get("total_records"), "0"))
    add_textbox(project, Inches(0.95), Inches(3.54), Inches(2.1), Inches(0.3), "数据覆盖", 12, NAVY, True)
    add_bullet_list(project, coverage, Inches(3.95), 5)

    data_slide = add_base_slide(presentation, "数据汇总", "仅展示已观测、可公开的数据")
    observed = as_mapping(summary.get("observed_metrics"), "observed_metrics")
    metric_cards = (
        ("播放量", metric_value(observed.get("views"))),
        ("互动量", metric_value(observed.get("interactions"))),
        ("互动率", metric_value(observed.get("engagement_rate"))),
        ("点赞数", metric_value(observed.get("likes"))),
        ("评论数", metric_value(observed.get("comments"))),
        ("收藏与转发", metric_value(observed.get("favorites")) + " / " + metric_value(observed.get("shares"))),
    )
    for index, (label, value) in enumerate(metric_cards):
        column = index % 3
        row = index // 3
        add_card(data_slide, Inches(0.93 + column * 4.0), Inches(1.55 + row * 2.0), Inches(3.7), Inches(1.52), label, value)

    comparison = add_base_slide(presentation, "平台与产品对比", "基于可比较内容的汇总表现")
    add_comparison_rows(comparison, as_mapping(sections["eligible_comparisons"], "eligible_comparisons"))

    indicator_slide = add_base_slide(presentation, "关键指标", "客户版不包含商业成本与投放归因数据")
    selected = as_mapping(indicators.get("selected_metric"), "selected_metric")
    add_card(indicator_slide, Inches(0.95), Inches(1.62), Inches(5.45), Inches(1.65), text(selected.get("label"), text(report.get("selected_metric"))), metric_value(selected), text(selected.get("definition"), ""))
    commercial = as_mapping(indicators.get("commercial"), "commercial")
    add_card(indicator_slide, Inches(6.9), Inches(1.62), Inches(5.45), Inches(1.65), "商业指标", text(commercial.get("disclosure"), "暂不包含"))
    add_textbox(indicator_slide, Inches(0.96), Inches(3.78), Inches(2.6), Inches(0.35), "指标说明", 12, NAVY, True)
    add_bullet_list(indicator_slide, [text(selected.get("definition"), "以客户已确认的数据口径为准。")], Inches(4.15), 2)

    cases_slide = add_base_slide(presentation, "优秀案例", "以可公开的内容摘要识别可复用方向")
    cases = as_mapping(sections["excellent_cases"], "excellent_cases")
    case_items = []
    for item in as_list(cases.get("cases")):
        if not isinstance(item, dict):
            continue
        case_items.append(" · ".join(filter(None, [text(item.get("reference"), ""), text(item.get("platform"), ""), metric_value(item.get("selected_metric"))])))
    add_bullet_list(cases_slide, case_items or [text(cases.get("reason"), "当前没有可公开的优秀案例。")], Inches(1.65), 7)

    limits_slide = add_base_slide(presentation, "数据边界与风险", "明确本次复盘可解释的数据范围")
    limitation_items = [text(item.get("disclosure") if isinstance(item, dict) else item, "") for item in as_list(limits.get("limitations"))]
    source_modes = [
        text(item.get("mode"), "已登记来源") + "：" + text(item.get("count"), "0") + " 条"
        for item in as_list(limits.get("source_modes")) if isinstance(item, dict)
    ]
    add_bullet_list(limits_slide, source_modes + limitation_items, Inches(1.62), 8)

    next_slide = add_base_slide(presentation, "优化建议与下一周期", "围绕当前数据证据推进下一轮内容优化")
    actions = as_list(optimization.get("optimization_actions"))
    add_textbox(next_slide, Inches(0.95), Inches(1.52), Inches(2.5), Inches(0.35), "优化建议", 12, NAVY, True)
    add_bullet_list(next_slide, actions, Inches(1.95), 5)
    add_textbox(next_slide, Inches(0.95), Inches(5.12), Inches(2.5), Inches(0.35), "下一周期计划", 12, NAVY, True)
    add_textbox(next_slide, Inches(0.95), Inches(5.55), Inches(11.1), Inches(0.85), text(optimization.get("next_cycle_plan"), "待与客户确认下一周期执行计划。"), 15, TEXT)
    return presentation


def generate(input_path, output_path):
    if not os.path.isabs(output_path) or not output_path.lower().endswith(".pptx"):
        abort("output path must be an absolute .pptx path")
    report, sections = load_report(input_path)
    presentation = build_presentation(report, sections)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    presentation.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        abort("usage: generate_customer_report_ppt.py INPUT_JSON OUTPUT_PPTX")
    generate(sys.argv[1], sys.argv[2])
